# Notes for frontend team :
# The topic trigger is removed from the pipeline (existes in /bin folder).
# this is a mathpix technique to extract the topics from the syllabus 
# api endpoints for it is /syllabus/extract

from fastapi import APIRouter, UploadFile, File, HTTPException
from fastapi.responses import JSONResponse
from typing import List
import io, json, re, base64, requests, os
from pathlib import Path
from dotenv import load_dotenv, find_dotenv
from pdf2image import convert_from_bytes
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook

router = APIRouter()
################################## CHANGE THE PAGE LIMIT TO SYLLABUS HERE
PAGE_LIMIT=20
##################################

# Load environment
BASE_DIR = Path(__file__).resolve().parent
load_dotenv(find_dotenv(), override=True)

MATHPIX_APP_ID = os.getenv("MATHPIX_APP_ID")
MATHPIX_API_KEY = os.getenv("MATHPIX_API_KEY")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

if not MATHPIX_APP_ID or not MATHPIX_API_KEY:
    raise RuntimeError("Mathpix credentials missing - check .env")

# Gemini optional
USE_GEMINI = bool(GEMINI_API_KEY)

try:
    from google import generativeai as genai
except Exception:
    genai = None

if USE_GEMINI and genai and GEMINI_API_KEY:
    genai.configure(api_key=GEMINI_API_KEY)



# Utilities
def encode_file(file_bytes: bytes, content_type: str) -> str:
    b64 = base64.b64encode(file_bytes).decode("utf-8")
    return f"data:{content_type};base64,{b64}"


def call_mathpix_api(data: dict) -> dict:
    headers = {
        "app_id": MATHPIX_APP_ID,
        "app_key": MATHPIX_API_KEY,
        "Content-type": "application/json",
    }
    resp = requests.post("https://api.mathpix.com/v3/text", headers=headers, json=data, timeout=120)

    try:
        result = resp.json()
    except Exception:
        raise HTTPException(status_code=resp.status_code, detail="Mathpix returned non-JSON")

    if resp.status_code != 200 or "error" in result:
        raise HTTPException(status_code=resp.status_code, detail=result)

    return result



# File extractors
def extract_text_from_docx(b: bytes) -> str:
    doc = Document(io.BytesIO(b))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_text_from_pptx(b: bytes) -> str:
    prs = Presentation(io.BytesIO(b))
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                texts.append(shape.text)
    return "\n".join(texts)


def extract_text_from_xlsx(b: bytes) -> str:
    wb = load_workbook(io.BytesIO(b))
    texts = []
    for sheet in wb:
        for row in sheet.iter_rows(values_only=True):
            texts.append(" ".join(str(cell) for cell in row if cell is not None))
    return "\n".join(texts)



# Gemini topic extractor
async def extract_topics_with_gemini(text: str, subject: str = "General") -> list:
    if not genai:
        raise HTTPException(500, "Gemini not available")

    prompt = f"""You're an expert at identifying academic topics in a syllabus or textbook.
    Directly take the topics as they are written in the file(s). The topic names should exactly match the text.
    Extract a list of academic {subject} topics from the provided text.

    Return a JSON array of objects:
      - "topic": the academic topic (concise but complete)
      - "summary": a brief explanation of what that topic covers

    Guidelines:
    - Only include real academic topics.
    - Avoid duplicates or vague items.
    - Keep names short and clear.
    - Return raw JSON only, no markdown or prose.
    """

    model = genai.GenerativeModel("gemini-2.5-flash")

    response = model.generate_content(prompt + "\n\n" + text[:30000])

    cleaned = re.sub(r"```json|```", "", response.text).strip()
    match = re.search(r"\[.*\]", cleaned, re.S)

    if match:
        try:
            topics = json.loads(match.group(0))
            # Normalize whitespace
            for t in topics:
                if "topic" in t:
                    t["topic"] = re.sub(r"\s+", " ", t["topic"]).strip()
                if "summary" in t:
                    t["summary"] = re.sub(r"\s+", " ", t["summary"]).strip()
            return topics
        except:
            return cleaned

    return cleaned



# ROUTE
@router.post("/extract")
async def extract_text_and_topics(
    files: List[UploadFile] = File(...),
    subject: str = "General"
):
    combined_text = ""
    processed_files = []
    total_chars = 0

    try:
        for file in files:
            file_bytes = await file.read()
            filename = file.filename.lower()
            ext = Path(filename).suffix

            if ext == ".pdf":
                pages = convert_from_bytes(file_bytes, dpi=150)  # dpi low = faster
                num_pages = len(pages)

                if num_pages > PAGE_LIMIT:
                    raise HTTPException(
                        status_code=400,
                        detail=f"PDF '{filename}' has {num_pages} pages. Maximum allowed is {PAGE_LIMIT}."
                    )
                
                
            elif ext == ".pptx":
                prs = Presentation(io.BytesIO(file_bytes))
                if len(prs.slides) > PAGE_LIMIT:
                    raise HTTPException(400, f"PPTX '{filename}' has more than {PAGE_LIMIT} slides.")
            elif ext == ".docx":
                DOCX_PARAGRAPHS_PER_PAGE = 15
                doc = Document(io.BytesIO(file_bytes))
                paragraph_count = len(doc.paragraphs)
                max_paragraphs = PAGE_LIMIT * DOCX_PARAGRAPHS_PER_PAGE
                if paragraph_count > max_paragraphs:
                    raise HTTPException(
                        status_code=400,
                        detail=(
                            f"DOCX '{filename}' is too long. "
                            f"Detected approximately {paragraph_count // DOCX_PARAGRAPHS_PER_PAGE} pages "
                            f"(limit is {PAGE_LIMIT})."
                        )
                    )

            file_text = ""

            
            # PDF -> Mathpix OCR
            if ext == ".pdf":
                pages = convert_from_bytes(file_bytes, dpi=200)
                for i, img in enumerate(pages, 1):
                    buf = io.BytesIO()
                    img.save(buf, format="PNG")
                    encoded = encode_file(buf.getvalue(), "image/png")

                    result = call_mathpix_api({
                        "src": encoded,
                        "formats": ["text"],
                        "ocr": ["math", "text"]
                    })

                    file_text += f"\n--- Page {i} ---\n{result.get('text', '')}"

            
            # Image -> Mathpix OCR
            elif ext in [".png", ".jpg", ".jpeg", ".bmp", ".tiff"]:
                encoded = encode_file(file_bytes, f"image/{ext.lstrip('.')}")
                result = call_mathpix_api({
                    "src": encoded,
                    "formats": ["text"],
                    "ocr": ["math", "text"]
                })
                file_text = result.get("text", "")

            
            # Text formats
            elif ext in [".txt", ".md", ".csv", ".json"]:
                text = file_bytes.decode("utf-8", errors="ignore")
                file_text = json.dumps(json.loads(text), indent=2) if ext == ".json" else text

            
            # Office files
            elif ext == ".docx":
                file_text = extract_text_from_docx(file_bytes)

            elif ext == ".pptx":
                file_text = extract_text_from_pptx(file_bytes)

            elif ext == ".xlsx":
                file_text = extract_text_from_xlsx(file_bytes)

            else:
                # fallback - raw text
                file_text = file_bytes.decode(errors="ignore")

            # Collect
            chars = len(file_text)
            total_chars += chars
            combined_text += f"\n\n===== FILE: {filename} =====\n{file_text}"

            processed_files.append({
                "filename": filename,
                "characters": chars
            })

        if total_chars == 0:
            raise HTTPException(400, "No readable text extracted.")

        # Gemini topics
        topics = await extract_topics_with_gemini(combined_text, subject) if USE_GEMINI else []

        return JSONResponse(content={
            "success": True,
            "file_count": len(processed_files),
            "processed_files": processed_files,
            "total_characters": total_chars,
            "topics": topics,
            "full_text": combined_text
        })

    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(500, f"Internal server error: {e}")



# OPTIONAL health routes
@router.get("/health")
def health_check():
    return {"status": "healthy", "mathpix_configured": bool(MATHPIX_APP_ID and MATHPIX_API_KEY)}
